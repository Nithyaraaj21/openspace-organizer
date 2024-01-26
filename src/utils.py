from openpyxl import Workbook

class Store:
    def __init__(self, tables):
        self.tables = tables

    def store(self, filename):
        wb = Workbook()
        for i, table in enumerate(self.tables):
            ws_table = wb.create_sheet(f"Table {i+1}")
            ws_table.append(['Name'])
            for seat in table.seats:
                if not seat.free:
                    ws_table.append([str(seat.seat_occupant)])
        wb.remove(wb['Sheet'])
        wb.save(filename)